"""Generate the example planning sheet shipped with the repo.

The example is deliberately untidy - a title row above the headers, three
different date formats, one cell holding a whole time range, a session with
only a duration, an all-day event, Danish column names, and one row that
cannot be converted - so that running the converter on it demonstrates what
it tolerates and what it reports.

    python examples/make_example_sheet.py
"""

from __future__ import annotations

import csv
import datetime as dt
from pathlib import Path

HERE = Path(__file__).parent

HEADER = [
    "Date",
    "Time",
    "End time",
    "Duration",
    "Subject",
    "Location",
    "Content",
    "Underviser",
    "Track",
]

ROWS = [
    [
        dt.date(2026, 9, 7), "09:00", "12:00", None,
        "Kick-off: What is entrepreneurship at DTU?",
        "Building 358, Aud. 060",
        "Welcome, programme overview, team formation exercise.",
        "Toke Reichstein", "Core",
    ],
    [
        "08/09/2026", "13:00-16:00", None, None,
        "Opportunity recognition",
        "Building 421, Room 073",
        "Idea generation, effectuation, and the customer problem.",
        "Toke Reichstein", "Core",
    ],
    [
        "9. september 2026", "9", None, "90",
        "Guest lecture: from lab to launch",
        "Skylab, Innovation Hub",
        "A DTU spin-out founder on the first 18 months.",
        "External speaker", "Guest",
    ],
    [
        dt.date(2026, 9, 10), None, None, None,
        "Individual supervision - book a slot",
        "Skylab",
        "Sign-up sheet circulated the week before.",
        "Teaching assistants", "Supervision",
    ],
    [
        dt.date(2026, 9, 11), "10:00", "15:00", None,
        "Business model workshop",
        "Skylab, Workshop Room 2",
        "Business Model Canvas applied to your own case; peer feedback.",
        "Toke Reichstein", "Workshop",
    ],
    [
        dt.date(2026, 9, 14), "16:00", "19:00", None,
        "Pitch training",
        "Building 101, Meeting Room A",
        "Structure, delivery, and handling the hard questions.",
        "Student assistants", "Workshop",
    ],
    [
        dt.date(2026, 9, 18), "13:00", "17:00", None,
        "Final pitch and feedback panel",
        "Skylab, Main Hall",
        "Ten-minute pitches to a panel of investors and DTU staff.",
        "Toke Reichstein", "Assessment",
    ],
    [
        # A row that cannot be converted - it is reported, not silently dropped.
        "TBD", "09:00", "12:00", None,
        "Follow-up session (date not fixed yet)",
        "",
        "Placeholder while we wait for the room booking.",
        "", "Core",
    ],
]


def write_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    book = Workbook()
    sheet = book.active
    sheet.title = "Programme"

    sheet.append(["DTU Entrepreneurship - Autumn 2026 programme (draft)"])
    sheet.append([])
    sheet.append(HEADER)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    for row in ROWS:
        sheet.append(row)

    widths = [14, 14, 10, 10, 44, 26, 52, 20, 14]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[sheet.cell(row=3, column=index).column_letter].width = width
    for row in sheet.iter_rows(min_row=4, max_col=1):
        for cell in row:
            if isinstance(cell.value, dt.date):
                cell.number_format = "dd/mm/yyyy"

    book.save(path)


def write_csv(path: Path) -> None:
    with open(path, "w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["DTU Entrepreneurship - Autumn 2026 programme (draft)"])
        writer.writerow([])
        writer.writerow(HEADER)
        for row in ROWS:
            writer.writerow(
                [
                    value.isoformat() if isinstance(value, dt.date) else (value or "")
                    for value in row
                ]
            )


if __name__ == "__main__":
    write_xlsx(HERE / "programme_autumn2026.xlsx")
    write_csv(HERE / "programme_autumn2026.csv")
    print("wrote programme_autumn2026.xlsx and .csv in", HERE)
