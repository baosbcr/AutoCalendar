"""Reading the planning spreadsheet.

Nobody is going to rewrite their sheet to match our column names, so the
reader recognises the headers people actually use - in English and Danish -
finds the header row wherever it sits, and keeps any column it does not
recognise by appending it to the event description. Rows it cannot use are
reported rather than dropped silently.
"""

from __future__ import annotations

import csv
import datetime as dt
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path

from .model import Event
from .parsing import (
    parse_bool,
    parse_date,
    parse_duration,
    parse_time,
    split_time_range,
)

# Canonical field -> the header spellings that mean it. Compared after
# lowercasing and stripping everything that is not a letter or digit.
ALIASES: dict[str, tuple[str, ...]] = {
    "title": (
        "title", "subject", "subjectname", "subjects", "name", "session",
        "topic", "activity", "event", "course", "module", "headline",
        "emne", "emnenavn", "fag", "titel", "aktivitet", "kursus", "modul",
        "overskrift", "undervisning",
    ),
    "date": ("date", "startdate", "day", "when", "dato", "startdato", "dag"),
    "end_date": ("enddate", "todate", "slutdato", "tildato"),
    "start": (
        "start", "starttime", "from", "time", "timeslot", "begin", "starts",
        "kl", "klokken", "starttid", "tid", "tidspunkt", "fra", "fratid",
    ),
    "end": (
        "end", "endtime", "to", "until", "finish", "ends",
        "sluttid", "til", "tiltid",
    ),
    "duration": (
        "duration", "length", "minutes", "mins", "varighed", "laengde",
        "længde", "minutter", "varighedmin",
    ),
    "location": (
        "location", "room", "venue", "place", "building", "address", "where",
        "lokation", "lokale", "sted", "rum", "bygning", "adresse", "hvor",
    ),
    "description": (
        "description", "content", "notes", "details", "agenda", "comment",
        "comments", "summary", "info", "beskrivelse", "indhold", "noter",
        "detaljer", "kommentar", "kommentarer", "program",
    ),
    "url": ("url", "link", "links", "meetinglink", "teams", "zoom", "web"),
    "all_day": ("allday", "fullday", "heldag", "heledagen"),
    "category": ("category", "type", "track", "tag", "tags", "label", "kategori"),
    # Deliberately narrow. A column called "Speaker" or "Underviser" holds a
    # name, not a mailing list - treating it as one would invite people
    # nobody meant to invite. Only headers that unambiguously mean "these
    # people get an invitation" count.
    "required": (
        "required", "attendees", "attendee", "requiredattendees",
        "participants", "invitees", "deltagere", "inviterede",
    ),
    "optional": (
        "optional", "optionalattendees", "cc", "copy", "observers",
        "valgfri", "valgfrie", "kopi",
    ),
    "calendar": ("calendar", "targetcalendar", "kalender", "calendarname"),
}

_CANONICAL_BY_ALIAS = {
    alias: canonical for canonical, names in ALIASES.items() for alias in names
}

DEFAULT_DURATION = dt.timedelta(hours=1)


@dataclass
class RowProblem:
    """Something that stopped, or bent, a row on its way to becoming an event."""

    row: int
    message: str
    fatal: bool = True  # fatal rows produce no event

    def __str__(self) -> str:
        kind = "skipped" if self.fatal else "warning"
        return f"row {self.row}: {kind} - {self.message}"


@dataclass
class ReadResult:
    events: list[Event] = field(default_factory=list)
    problems: list[RowProblem] = field(default_factory=list)
    mapping: dict[str, str] = field(default_factory=dict)  # canonical -> header
    extra_columns: list[str] = field(default_factory=list)
    header_row: int = 1
    sheet: str = ""

    @property
    def skipped(self) -> list[RowProblem]:
        return [p for p in self.problems if p.fatal]

    @property
    def warnings(self) -> list[RowProblem]:
        return [p for p in self.problems if not p.fatal]


def normalise(header) -> str:
    return "".join(ch for ch in str(header or "").lower() if ch.isalnum())


def _is_blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _text(value) -> str:
    if _is_blank(value):
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


# --------------------------------------------------------------------------
# Loading raw rows
# --------------------------------------------------------------------------

def _load_xlsx(path: Path, sheet: str | int | None) -> tuple[list[list], str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on the environment
        raise RuntimeError(
            "Reading .xlsx needs openpyxl. Install it with: pip install openpyxl\n"
            "(or save the sheet as CSV, which needs nothing extra)"
        ) from exc

    book = load_workbook(path, data_only=True, read_only=True)
    try:
        if sheet is None:
            worksheet = book.active
        elif isinstance(sheet, int):
            worksheet = book.worksheets[sheet]
        else:
            if sheet not in book.sheetnames:
                raise ValueError(
                    f"no sheet named {sheet!r}; the file has: "
                    + ", ".join(book.sheetnames)
                )
            worksheet = book[sheet]
        rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
        name = worksheet.title
    finally:
        book.close()
    return rows, name


def _load_csv(path: Path) -> tuple[list[list], str]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover - latin-1 decodes anything
        raise RuntimeError(f"could not decode {path}")

    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel()
        dialect.delimiter = ";" if sample.count(";") > sample.count(",") else ","
    rows = [list(row) for row in csv.reader(text.splitlines(), dialect)]
    return rows, path.stem


def split_people(text: str) -> list[str]:
    """Split an attendee cell into names or addresses.

    Semicolons win when present, because that is what Outlook writes and
    because a display name may legitimately contain a comma ("Nielsen, Jan").
    Only when there is no semicolon at all do commas separate.
    """
    if not text:
        return []
    separator = ";" if ";" in text else ","
    parts = (piece.strip() for chunk in text.splitlines() for piece in chunk.split(separator))
    return [p for p in parts if p]


def load_rows(path: Path, sheet: str | int | None = None) -> tuple[list[list], str]:
    """Return the raw grid of cells, plus a name for what was read."""
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return _load_xlsx(path, sheet)
    if suffix in (".csv", ".tsv", ".txt"):
        return _load_csv(path)
    if suffix == ".xls":
        raise ValueError(
            ".xls is the old Excel format - open it in Excel and save as .xlsx or .csv"
        )
    raise ValueError(f"unsupported file type {suffix!r}; use .xlsx or .csv")


# --------------------------------------------------------------------------
# Finding the header row and mapping the columns
# --------------------------------------------------------------------------

def find_header_row(rows: Sequence[Sequence], look_ahead: int = 15) -> int:
    """Index of the most header-like row within the first ``look_ahead`` rows.

    Planning sheets often open with a title and a blank line, so the headers
    are rarely on row 1.
    """
    best_index, best_score = 0, 0
    for index, row in enumerate(rows[:look_ahead]):
        canon = {_CANONICAL_BY_ALIAS.get(normalise(cell)) for cell in row}
        score = len(canon - {None})
        if {"title", "date"} <= canon:  # a believable header has both
            score += 2
        if score > best_score:
            best_index, best_score = index, score
    return best_index


def map_columns(header: Sequence) -> tuple[dict[str, int], list[tuple[int, str]]]:
    """Map canonical field -> column index, plus the columns we did not claim."""
    mapping: dict[str, int] = {}
    extras: list[tuple[int, str]] = []
    for index, cell in enumerate(header):
        label = _text(cell)
        if not label:
            continue
        canonical = _CANONICAL_BY_ALIAS.get(normalise(cell))
        if canonical and canonical not in mapping:
            mapping[canonical] = index
        else:
            extras.append((index, label))
    return mapping, extras


# --------------------------------------------------------------------------
# Rows -> events
# --------------------------------------------------------------------------

def _row_to_event(
    row: Sequence,
    row_number: int,
    columns: dict[str, int],
    extras: Sequence[tuple[int, str]],
    default_duration: dt.timedelta,
    problems: list[RowProblem],
) -> Event | None:
    def cell(field_name: str):
        index = columns.get(field_name)
        if index is None or index >= len(row):
            return None
        return row[index]

    title = _text(cell("title"))
    if not title:
        problems.append(RowProblem(row_number, "no subject/title in the row"))
        return None

    if _is_blank(cell("date")):
        problems.append(RowProblem(row_number, f"no date for {title!r}"))
        return None
    try:
        day = parse_date(cell("date"))
    except ValueError as exc:
        problems.append(RowProblem(row_number, f"{exc} (for {title!r})"))
        return None

    end_day = day
    if not _is_blank(cell("end_date")):
        try:
            end_day = max(parse_date(cell("end_date")), day)
        except ValueError as exc:
            problems.append(RowProblem(row_number, f"end date: {exc}", fatal=False))

    # A start cell may itself hold a whole range: "09:00-12:00".
    start_text, embedded_end = split_time_range(cell("start"))
    end_raw = cell("end")
    all_day_flag = (
        parse_bool(cell("all_day")) if columns.get("all_day") is not None else False
    )

    if all_day_flag or (not start_text and _is_blank(end_raw)):
        event_start: dt.date = day
        event_end: dt.date = end_day
    else:
        if start_text:
            try:
                start_time = parse_time(start_text)
            except ValueError as exc:
                problems.append(
                    RowProblem(row_number, f"start time: {exc} (for {title!r})")
                )
                return None
        else:
            start_time = dt.time(9, 0)
            problems.append(
                RowProblem(row_number, "no start time, assumed 09:00", fatal=False)
            )

        event_start = dt.datetime.combine(day, start_time)
        event_end = None
        end_source = end_raw if not _is_blank(end_raw) else embedded_end

        if not _is_blank(end_source):
            try:
                event_end = dt.datetime.combine(end_day, parse_time(end_source))
            except ValueError as exc:
                problems.append(RowProblem(row_number, f"end time: {exc}", fatal=False))

        if event_end is None and not _is_blank(cell("duration")):
            try:
                event_end = event_start + parse_duration(cell("duration"))
            except ValueError as exc:
                problems.append(RowProblem(row_number, f"duration: {exc}", fatal=False))

        if event_end is None:
            event_end = event_start + default_duration
        elif event_end < event_start:  # an evening session running past midnight
            event_end += dt.timedelta(days=1)
            problems.append(
                RowProblem(
                    row_number,
                    "end is before start, assumed it runs past midnight",
                    fatal=False,
                )
            )
        elif event_end == event_start:
            event_end = event_start + default_duration

    description_parts = []
    if _text(cell("description")):
        description_parts.append(_text(cell("description")))
    for index, label in extras:
        if index < len(row) and not _is_blank(row[index]):
            description_parts.append(f"{label}: {_text(row[index])}")

    categories = [
        part.strip()
        for part in _text(cell("category")).replace(";", ",").split(",")
        if part.strip()
    ]

    return Event(
        title=title,
        start=event_start,
        end=event_end,
        location=_text(cell("location")),
        description="\n".join(description_parts),
        categories=categories,
        url=_text(cell("url")),
        required=split_people(_text(cell("required"))),
        optional=split_people(_text(cell("optional"))),
        calendar=_text(cell("calendar")),
        source_row=row_number,
    )


def read_events(
    path,
    *,
    sheet: str | int | None = None,
    default_duration: dt.timedelta = DEFAULT_DURATION,
    keep_extra_columns: bool = True,
) -> ReadResult:
    """Read a spreadsheet into events, collecting problems as it goes."""
    path = Path(path)
    rows, sheet_name = load_rows(path, sheet)
    result = ReadResult(sheet=sheet_name)
    if not rows:
        result.problems.append(RowProblem(0, "the file is empty"))
        return result

    header_index = find_header_row(rows)
    result.header_row = header_index + 1
    columns, extras = map_columns(rows[header_index])
    result.mapping = {
        canonical: _text(rows[header_index][index])
        for canonical, index in columns.items()
    }
    result.extra_columns = [label for _, label in extras]

    missing = {"title", "date"} - set(columns)
    if missing:
        found = ", ".join(sorted(result.mapping.values())) or "nothing recognisable"
        result.problems.append(
            RowProblem(
                result.header_row,
                f"could not find a {' and a '.join(sorted(missing))} column "
                f"(columns recognised on the header row: {found})",
            )
        )
        return result

    if not keep_extra_columns:
        extras = []

    for row_number, row in enumerate(
        rows[header_index + 1 :], start=header_index + 2
    ):
        if all(_is_blank(cell) for cell in row):
            continue
        event = _row_to_event(
            row, row_number, columns, extras, default_duration, result.problems
        )
        if event is not None:
            result.events.append(event)

    result.events.sort(key=lambda e: (str(e.start), e.title))
    return result
